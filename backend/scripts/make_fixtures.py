"""Write small synthetic documents for manual integration; no real business data."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas

TARGET = Path(__file__).resolve().parents[1] / "tests/fixtures"


def main():
    TARGET.mkdir(parents=True, exist_ok=True)
    for version, lines in {
        "before": [
            "1.1. The Audit department prepares a quarterly risk report.",
            "1.2. The Audit department archives audit working papers for five years.",
        ],
        "after": [
            "1.1. The Risk department prepares a quarterly risk report.",
            "1.2. The Risk department monitors remediation deadlines.",
        ],
    }.items():
        document = Document()
        document.add_heading("SYNTHETIC TEST / " + version.upper(), 0)
        document.add_heading("Audit department" if version == "before" else "Risk department", 1)
        for line in lines:
            document.add_paragraph(line)
        document.save(TARGET / f"{version}.docx")
        canvas = Canvas(str(TARGET / f"{version}.pdf"), invariant=1)
        canvas.drawString(40, 780, "SYNTHETIC TEST / " + version.upper())
        for i, line in enumerate(lines):
            canvas.drawString(40, 740 - i * 30, line)
        canvas.save()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = version
        sheet.append(["SYNTHETIC TEST", "Function"])
        for i, line in enumerate(lines, 1):
            sheet.append([i, line])
        workbook.save(TARGET / f"{version}.xlsx")


if __name__ == "__main__":
    main()
