"""Reproducible Kazakh synthetic documents for the judges' acceptance scenario."""

from pathlib import Path

from docx import Document
from docx.shared import Pt

TARGET = Path(__file__).resolve().parents[1] / "tests/fixtures/judge"
REGISTER = "ұйымның барлық келісімшарттарын бірыңғай тізілімге тіркеуге міндетті."
ARCHIVE = "ұйымның барлық келісімшарттарының түпнұсқаларын орталық мұрағатта сақтауға міндетті."
QUALITY = "Сапа бөлімі қызмет көрсету жөніндегі шағымдарды тоқсан сайын талдауға міндетті."


def create(version, sections):
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("Ұйымдық өзгерістер — бақылау құжаты", 0)
    document.add_paragraph("ЖАСАНДЫ ДЕРЕКТЕР. Нақты ұйымға немесе қызметкерге қатысты емес.")
    document.add_paragraph("Редакция: " + ("дейін" if version == "before" else "кейін"))
    for heading, clauses in sections:
        document.add_heading(heading, 1)
        for clause in clauses:
            document.add_paragraph(clause)
    document.save(TARGET / f"{version}.docx")


def main():
    TARGET.mkdir(parents=True, exist_ok=True)
    create(
        "before",
        [
            ("1. Ұйымдық құрылым", ["1.1. Ұйым құрамында Құжат айналымы бөлімі және Сапа бөлімі бар."]),
            (
                "2. Құжат айналымы бөлімі",
                [
                    "2.1. Құжат айналымы бөлімі " + REGISTER,
                    "2.2. Құжат айналымы бөлімі электрондық құжаттардың резервтік көшірмелерін әр апта сайын тексеруге міндетті.",
                ],
            ),
            ("3. Сапа бөлімі", ["3.1. " + QUALITY]),
        ],
    )
    create(
        "after",
        [
            (
                "1. Қайта ұйымдастыру туралы өкім",
                [
                    "1.1. Құжат айналымы бөлімі Ақпаратты басқару бөлімі болып қайта аталсын.",
                    "1.2. Сапа бөлімі сақталсын. Құжаттарды сақтау бөлімі құрылсын.",
                ],
            ),
            (
                "2. Ақпаратты басқару бөлімі",
                ["2.1. Ақпаратты басқару бөлімі " + REGISTER, "2.2. Ақпаратты басқару бөлімі " + ARCHIVE],
            ),
            ("3. Сапа бөлімі", ["3.1. " + QUALITY]),
            ("4. Құжаттарды сақтау бөлімі", ["4.1. Құжаттарды сақтау бөлімі " + ARCHIVE]),
        ],
    )
    print("Created two synthetic judge documents; no API calls.")


if __name__ == "__main__":
    main()
