import io
import zipfile
from dataclasses import replace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas

from backend.app.models import ParsedDocument
from backend.app.parsers import RejectedFile, inspect_archive, parse_document
from backend.tests.conftest import docx_bytes


def save(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def parse(content, kind, settings):
    result = parse_document(content, "test." + kind, kind, "doc-test", "before", settings)
    ParsedDocument.model_validate(result)
    return result


def test_docx_clauses_table_context_and_raw(settings):
    doc = Document()
    doc.add_heading("5. Chief auditor must not:", 1)
    raw = "3.9 б. Руководитель. 3.10.Рабочие места. 3.11.Работники могут работать. 3.12.По вопросам работы."
    doc.add_paragraph(raw)
    doc.add_paragraph("5.5.3 ;")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "6.1. Annual report"
    table.cell(0, 1).text = "Audit department"
    doc.sections[0].header.paragraphs[0].text = "Repeated header"
    result = parse(save(doc), "docx", settings)
    spans = result["spans"]
    assert {"3.9", "3.10", "3.11", "3.12", "5.5.3", "6.1"} <= {s["locator"]["clause"] for s in spans}
    parts = [s for s in spans if "/p[2]/" in s["locator"]["path"] and "/tbl[" not in s["locator"]["path"]]
    assert "".join(s["raw_text"] for s in parts) == raw
    assert all(s["locator"]["page"] is None for s in spans)
    empty = next(s for s in spans if s["locator"]["clause"] == "5.5.3")
    assert not empty["is_content"]
    assert any(w["code"] == "empty_clause" for w in result["warnings"])
    assert not next(s for s in spans if s["raw_text"] == "Repeated header")["is_content"]
    assert any("/tbl[1]/tr[1]/tc[2]/" in s["locator"]["path"] for s in spans)
    assert all(spans[0]["id"] in s["context_span_ids"] for s in parts)


def test_docx_does_not_split_inline_references(settings):
    raw = "1.1. Comply with clause 3.4 and section 7.2 of the policy."
    result = parse(docx_bytes(raw), "docx", settings)
    assert len([s for s in result["spans"] if s["is_content"]]) == 1
    assert result["spans"][1]["raw_text"] == raw


def test_numbering_styles_and_source_stability(settings):
    doc = Document()
    doc.add_paragraph("First obligation", "List Number")
    doc.add_paragraph("Second obligation", "List Number")
    data = save(doc)
    first, second = parse(data, "docx", settings), parse(data, "docx", settings)
    assert [s["locator"]["clause"] for s in first["spans"]] == ["1", "2"]
    assert first == second
    other = parse_document(data, "test.docx", "docx", "doc-other", "before", settings)
    assert first["spans"][0]["id"] != other["spans"][0]["id"]
    assert first["spans"][0]["raw_text"] == "First obligation"


def test_toc_excluded(settings):
    doc = Document()
    paragraph = doc.add_paragraph("Annual report ............ 3")
    props = paragraph._p.get_or_add_pPr()
    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), "TOC1")
    props.append(style)
    doc.add_paragraph("1.1. Prepare the annual report.")
    result = parse(save(doc), "docx", settings)
    assert not result["spans"][0]["is_content"]
    assert result["spans"][1]["is_content"]


def test_text_limit_is_explicit(settings):
    result = parse(docx_bytes("1.1. " + "X" * 200), "docx", replace(settings, max_text_chars=40))
    assert result["parse_status"] == "failed"
    assert any(w["code"] == "text_limit" for w in result["warnings"])


def test_xlsx_sheet_coordinates_formulas_and_multisheet(settings):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Құрылым"
    sheet["A1"] = "Audit department"
    sheet["B2"] = "Annual report"
    sheet["C3"] = '=HYPERLINK("https://invalid.test", "link")'
    workbook.create_sheet("After")["D4"] = "New task"
    result = parse(save(workbook), "xlsx", settings)
    assert result["parse_status"] == "partial"
    assert len(result["spans"]) == 4
    assert {s["locator"]["cell_range"] for s in result["spans"]} == {"A1", "B2", "C3", "D4"}
    formula = next(s for s in result["spans"] if s["locator"]["cell_range"] == "C3")
    assert not formula["is_content"] and formula["raw_text"].startswith("=HYPERLINK")
    assert any(w["code"] == "formula_not_evaluated" for w in result["warnings"])


def test_pdf_page_numbers_and_empty_page(settings):
    output = io.BytesIO()
    canvas = Canvas(output)
    canvas.drawString(30, 700, "1.1. Annual audit report")
    canvas.showPage()
    canvas.showPage()
    canvas.save()
    result = parse(output.getvalue(), "pdf", settings)
    assert result["parse_status"] == "partial"
    assert result["spans"][0]["locator"]["page"] == 1
    assert any(w["code"] == "page_without_text" for w in result["warnings"])


def test_encrypted_pdf_rejected(settings):
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(RejectedFile, match="encrypted_pdf"):
        parse(output.getvalue(), "pdf", settings)


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pdf"])
def test_corrupt_file_rejected(settings, kind):
    with pytest.raises(RejectedFile):
        parse(b"bad data", kind, settings)


def test_zip_bomb_and_entity_rejected(settings):
    with pytest.raises(RejectedFile, match="archive_limit"):
        inspect_archive(docx_bytes(), "docx", 10)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as z:
        z.writestr("[Content_Types].xml", "<a/>")
        z.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY boom SYSTEM "file:///etc/passwd">]><x>&boom;</x>')
    with pytest.raises(RejectedFile, match="invalid_container"):
        inspect_archive(output.getvalue(), "docx", settings.max_archive_bytes)
