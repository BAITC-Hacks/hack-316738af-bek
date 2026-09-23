import asyncio
import io
import time
import uuid
from copy import deepcopy
from functools import lru_cache

import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.app.config import Settings
from backend.app.main import create_app
from backend.app.parsers import parse_document
from backend.app.validation import evidence_references


@lru_cache(maxsize=32)
def docx_bytes(text="1.1. The audit department prepares an annual report."):
    # Repeated uploads must reuse identical bytes, including ZIP timestamps.
    document = Document()
    document.add_heading("Audit department", level=1)
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def result_for(payload):
    """A deliberately simple adapter fixture. It claims no model-quality evaluation."""
    docs = payload["documents"]
    sources = {
        v: next(s["id"] for d in docs if d["version"] == v for s in d["spans"] if s["is_content"])
        for v in ("before", "after")
    }
    units = [
        {
            "id": "u-" + v,
            "version": v,
            "name": "Audit",
            "aliases": [],
            "kind": "department",
            "parent_id": None,
            "source_span_ids": [sources[v]],
        }
        for v in sources
    ]
    functions = [
        {
            "id": "f-" + v,
            "version": v,
            "unit_id": "u-" + v,
            "role_id": None,
            "action": "report",
            "object": "audit",
            "scope": None,
            "modality": "obligation",
            "conditions": [],
            "frequency": None,
            "deliverable": None,
            "recipient": None,
            "source_span_ids": [sources[v]],
            "context_span_ids": [],
        }
        for v in sources
    ]
    result = {
        "schema_version": "1.0.0",
        "analysis_id": payload["analysis_id"],
        "analysis_revision": payload["analysis_revision"],
        "status": "partial" if any(d["parse_status"] != "complete" for d in docs) else "completed",
        "documents": [{k: deepcopy(v) for k, v in d.items() if k != "spans"} for d in docs],
        "units": units,
        "unit_changes": [],
        "functions": functions,
        "mappings": [
            {
                "id": "m1",
                "before_function_ids": ["f-before"],
                "after_function_ids": ["f-after"],
                "coverage_status": "full",
                "change_flags": [],
                "uncovered_aspects": [],
                "explanation": "Synthetic adapter fixture",
                "source_span_ids": list(sources.values()),
                "context_evidence_ids": [],
            }
        ],
        "findings": [
            {
                "id": "finding-1",
                "analysis_revision": payload["analysis_revision"],
                "type": "document_quality",
                "risk_change": "unknown",
                "title": "Synthetic check",
                "explanation": "A test, not a real AI finding",
                "severity": "low",
                "evidence_ids": [sources["before"]],
                "context_evidence_ids": [],
                "counterevidence_ids": [],
                "affected_units": ["u-before"],
                "recommended_action": {
                    "action": "Review source",
                    "target_role_id": None,
                    "required_document": None,
                    "reason": "Test fixture",
                },
                "verification_status": "needs_review",
                "review_status": "unreviewed",
                "search_scope_document_ids": [],
                "search_complete": False,
            }
        ],
        "coverage": {
            "files_total": len(docs),
            "files_complete": sum(d["parse_status"] == "complete" for d in docs),
            "files_partial": sum(d["parse_status"] == "partial" for d in docs),
            "files_failed": sum(d["parse_status"] == "failed" for d in docs),
            "before_functions_total": 1,
            "full": 1,
            "partial": 0,
            "none": 0,
            "unknown": 0,
            "after_functions_new": 0,
            "evidence_links_total": 0,
            "evidence_links_valid": 0,
        },
        "summary": {
            "headline": "Synthetic test report",
            "items": [{"text": "A test mapping", "reference_ids": ["m1"]}],
        },
        "usage": [],
        "warnings": [],
    }
    count = len(list(evidence_references(result)))
    result["coverage"]["evidence_links_total"] = result["coverage"]["evidence_links_valid"] = count
    return result


async def fake_engine(payload, emit_progress=None):
    if emit_progress:
        await emit_progress({"stage": "matching", "processed": 1, "total": 2, "message": "must not leak"})
    await asyncio.sleep(0.01)
    return result_for(payload)


@pytest.fixture
def settings(tmp_path):
    return Settings(data_dir=tmp_path / "data", frontend_dir=tmp_path / "frontend", analysis_timeout=2)


@pytest.fixture
def app(settings):
    return create_app(settings, engine=fake_engine, parser=parse_document)


@pytest.fixture
def client(app):
    with TestClient(app) as client:
        yield client


def upload(client, analysis_id, version, content=None, filename="audit.docx"):
    return client.post(
        f"/api/analyses/{analysis_id}/documents",
        data={"version": version},
        files={"file": (filename, content if content is not None else docx_bytes(), "application/octet-stream")},
    )


def prepared(client):
    state = client.post("/api/analyses").json()
    for v in ("before", "after"):
        response = upload(client, state["analysis_id"], v)
        assert response.status_code == 201, response.text
        state = response.json()
    return state


def start(client, state, key=None):
    return client.post(
        f"/api/analyses/{state['analysis_id']}/run",
        json={"expected_revision": state["analysis_revision"]},
        headers={"Idempotency-Key": key or str(uuid.uuid4())},
    )


def finish(client, state):
    for _ in range(150):
        response = client.get(f"/api/analyses/{state['analysis_id']}")
        current = response.json()
        if current["status"] in ("completed", "partial", "failed"):
            return current
        time.sleep(0.02)
    pytest.fail("Job did not finish")
