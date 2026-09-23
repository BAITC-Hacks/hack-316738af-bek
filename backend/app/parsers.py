"""Bounded document extraction. Raw evidence is preserved separately from labels.

No OCR, spreadsheet calculation, macros, network fetches or external links run here.
Parsers run in disposable processes (see workers.py).
"""

import hashlib
import io
import re
import zipfile
from collections import Counter

from defusedxml import ElementTree as SafeXML
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from pypdf import PdfReader

from .config import Settings
from .models import ParsedDocument

CLAUSE = re.compile(r"(?<!\S)(\d{1,3}(?:\.\d{1,3}){1,5})(?:[.)](?=\s|[^\W\d_])|(?=\s|$))")
LEADING = re.compile(
    r"^\s*((?:\d{1,3}(?:\.\d{1,3}){0,5})[.)]?|[а-яa-zәіңғүұқөһ][.)])(?:\s+|(?<=[.)])(?=[^\W\d_]))", re.I
)


class RejectedFile(ValueError):
    """Malformed or unsafe container, with a safe diagnostic code."""


def inspect_archive(content: bytes, kind: str, limit: int):
    """Inspect without extracting any paths to disk; reject XML entities and ZIP bombs."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [item.filename for item in entries]
            required = "word/document.xml" if kind == "docx" else "xl/workbook.xml"
            if required not in names or "[Content_Types].xml" not in names:
                raise RejectedFile("invalid_container")
            if len(entries) > 2500 or len(names) != len(set(names)):
                raise RejectedFile("archive_limit")
            if sum(item.file_size for item in entries) > limit:
                raise RejectedFile("archive_limit")
            flags = set()
            for item in entries:
                if item.flag_bits & 1 or item.file_size > limit:
                    raise RejectedFile("encrypted_or_oversized_archive")
                if "vbaProject" in item.filename:
                    raise RejectedFile("macros_not_supported")
                if any(s in item.filename for s in ("/media/", "/diagrams/", "/drawings/", "/embeddings/")):
                    flags.add("graphics_unread")
                if "externalLinks/" in item.filename:
                    flags.add("external_link_unread")
                if item.filename.endswith((".xml", ".rels")):
                    if item.file_size > 16 * 1024 * 1024:
                        raise RejectedFile("xml_limit")
                    SafeXML.fromstring(archive.read(item))
            return flags
    except RejectedFile:
        raise
    except Exception as exc:
        raise RejectedFile("invalid_container") from exc


class Builder:
    def __init__(self, document_id, kind, settings):
        self.id, self.kind, self.settings = document_id, kind, settings
        self.spans, self.warnings = [], []
        self.headings, self.clauses = {}, {}
        self.chars = 0
        self.stopped = False

    def warn(self, code, message, span_ids=None):
        value = {"code": code, "message": message, "document_id": self.id, "span_ids": span_ids or []}
        if value not in self.warnings:
            self.warnings.append(value)

    def add(self, raw, path, *, clause=None, page=None, sheet=None, cells=None, heading=None, content=True):
        if not raw.strip():
            return
        if self.stopped:
            return
        if self.chars + len(raw) > self.settings.max_text_chars or len(self.spans) >= self.settings.max_spans:
            self.stopped = True
            self.warn("text_limit", "Мәтін/үзінді шегінен кейінгі бөлік өңделмеді. Құжатты бөліп жүктеңіз.")
            return
        self.chars += len(raw)
        if heading is not None:
            self.headings = {k: v for k, v in self.headings.items() if k < heading}
            self.clauses.clear()
        leading = LEADING.match(raw)
        clause = clause or (leading.group(1).rstrip(".)") if leading else None)
        parents = list(self.headings.values())
        if clause:
            parts = clause.split(".")
            if parts[0].isdigit():
                self.clauses = {k: v for k, v in self.clauses.items() if clause.startswith(k + ".")}
                parents += [self.clauses[k] for k in self.clauses]
            elif self.clauses:
                parents += list(self.clauses.values())
        text_without_label = raw[leading.end() :] if leading else raw
        empty = not re.search(r"\w", text_without_label, re.UNICODE)
        span_id = "src_" + hashlib.sha256(f"{self.id}:{path}".encode()).hexdigest()[:28]
        span = {
            "id": span_id,
            "document_id": self.id,
            "raw_text": raw,
            "section_path": [p["raw_text"] for p in self.spans if p["id"] in self.headings.values()],
            "context_span_ids": list(dict.fromkeys(parents)),
            "locator": {
                "kind": self.kind,
                "clause": clause,
                "path": path,
                "page": page,
                "sheet": sheet,
                "cell_range": cells,
            },
            "is_content": content and heading is None and not empty,
        }
        self.spans.append(span)
        if heading is not None:
            self.headings[heading] = span_id
        if clause and clause[0].isdigit():
            self.clauses[clause] = span_id
        if empty:
            self.warn("empty_clause", "Тармақтың мәтіні бос немесе тек тыныс белгілерінен тұрады.", [span_id])

    def paragraph(self, raw, path, **kwargs):
        # Boundary offsets remain in the logical locator; raw slices are never normalized.
        matches = list(CLAUSE.finditer(raw))
        starts = sorted(
            set(
                [0]
                + [
                    m.start()
                    for m in matches
                    if m.start() > 0 and raw[: m.start()].rstrip().endswith((".", ";", ":", "!", "?"))
                ]
            )
        )
        for i, start in enumerate(starts):
            end = starts[i + 1] if i + 1 < len(starts) else len(raw)
            self.add(raw[start:end], f"{path}/text[{start}:{end}]", **kwargs)


def _text(element):
    parts = []
    for node in element.iter():
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append("\t")
        elif node.tag in (qn("w:br"), qn("w:cr")):
            parts.append("\n")
    return "".join(parts)


class Numbering:
    def __init__(self, document):
        self.styles = {s.style_id: s for s in document.styles}
        self.counters = {}
        try:
            self.root = document.part.numbering_part.element
        except (KeyError, NotImplementedError):
            self.root = None

    def label(self, paragraph):
        props = paragraph.find(qn("w:pPr"))
        if props is None:
            return None
        num = props.find(qn("w:numPr"))
        style = props.find(qn("w:pStyle"))
        visited = set()
        current = self.styles.get(style.get(qn("w:val"))) if style is not None else None
        while num is None and current is not None and current.style_id not in visited:
            visited.add(current.style_id)
            ppr = current.element.find(qn("w:pPr"))
            num = ppr.find(qn("w:numPr")) if ppr is not None else None
            current = current.base_style
        if num is None or self.root is None:
            return None
        id_node, level_node = num.find(qn("w:numId")), num.find(qn("w:ilvl"))
        if id_node is None:
            return None
        num_id = id_node.get(qn("w:val"))
        if num_id == "0":
            return None
        level = int(level_node.get(qn("w:val"))) if level_node is not None else 0
        nodes = self.root.xpath(f'./w:num[@w:numId="{int(num_id)}"]')
        if not nodes:
            return None
        abstract_id = nodes[0].find(qn("w:abstractNumId")).get(qn("w:val"))
        abstract = self.root.xpath(f'./w:abstractNum[@w:abstractNumId="{int(abstract_id)}"]')
        if not abstract:
            return None
        levels = {int(n.get(qn("w:ilvl"))): n for n in abstract[0].findall(qn("w:lvl"))}
        for override in nodes[0].findall(qn("w:lvlOverride")):
            replacement = override.find(qn("w:lvl"))
            if replacement is not None:
                levels[int(override.get(qn("w:ilvl")))] = replacement
        spec = levels.get(level)
        if spec is None:
            return None
        fmt = spec.find(qn("w:numFmt"))
        if fmt is not None and fmt.get(qn("w:val")) == "bullet":
            return None
        counters = self.counters.setdefault(num_id, {})

        def initial(index):
            lvl = levels.get(index)
            start = lvl.find(qn("w:start")) if lvl is not None else None
            value = int(start.get(qn("w:val"))) if start is not None else 1
            for override in nodes[0].findall(qn("w:lvlOverride")):
                s = override.find(qn("w:startOverride"))
                if int(override.get(qn("w:ilvl"))) == index and s is not None:
                    value = int(s.get(qn("w:val")))
            return value

        counters[level] = counters.get(level, initial(level) - 1) + 1
        for index in list(counters):
            if index > level:
                del counters[index]
        template = spec.find(qn("w:lvlText"))
        label = template.get(qn("w:val")) if template is not None else f"%{level + 1}."
        for index in range(9):
            value = counters.get(index, initial(index))
            lvl = levels.get(index)
            f = lvl.find(qn("w:numFmt")) if lvl is not None else None
            style_name = f.get(qn("w:val")) if f is not None else "decimal"
            if style_name in ("lowerLetter", "upperLetter") and 1 <= value <= 26:
                rendered = chr((97 if style_name == "lowerLetter" else 65) + value - 1)
            else:
                rendered = str(value)
            label = label.replace(f"%{index + 1}", rendered)
        return label.rstrip(".)")


def parse_docx(content, builder):
    document = Document(io.BytesIO(content))
    numbering = Numbering(document)
    toc_depth = 0

    def walk(parent, path):
        nonlocal toc_depth
        counts = Counter()
        for element in parent:
            tag = element.tag.rsplit("}", 1)[-1]
            counts[tag] += 1
            child_path = f"{path}/{tag}[{counts[tag]}]"
            if tag == "p":
                raw = _text(element)
                props = element.find(qn("w:pPr"))
                style = props.find(qn("w:pStyle")) if props is not None else None
                style_id = style.get(qn("w:val"), "") if style is not None else ""
                instructions = " ".join(n.text or "" for n in element.iter(qn("w:instrText")))
                if "TOC " in instructions.upper():
                    toc_depth = 1
                toc = toc_depth > 0 or style_id.lower().startswith("toc")
                if toc_depth and any(n.get(qn("w:fldCharType")) == "end" for n in element.iter(qn("w:fldChar"))):
                    toc_depth = 0
                match = re.search(r"heading\s*(\d)", style_id, re.I)
                outline = props.find(qn("w:outlineLvl")) if props is not None else None
                heading = (
                    int(match.group(1))
                    if match
                    else (int(outline.get(qn("w:val"))) + 1 if outline is not None else None)
                )
                if element.xpath(".//w:drawing | .//w:pict | .//w:object | .//w:txbxContent"):
                    builder.warn(
                        "graphics_unread", "Сурет/SmartArt/мәтіндік блок толық танылмады; мәтіндік баламасы қажет."
                    )
                builder.paragraph(raw, child_path, clause=numbering.label(element), heading=heading, content=not toc)
            elif tag in ("tbl", "tr", "tc", "sdt", "sdtContent"):
                walk(element, child_path)
            elif tag == "altChunk":
                builder.warn("embedded_content_unread", "Кірістірілген сыртқы мазмұн оқылмады.")

    walk(document.element.body, "word/document.xml/body")
    # Headers are evidence, never additional obligations; identical section headers deduplicate.
    seen = set()
    for section_index, section in enumerate(document.sections, 1):
        for label, part in (("header", section.header), ("footer", section.footer)):
            for i, paragraph in enumerate(part.paragraphs, 1):
                text = paragraph.text
                if text.strip() and text not in seen:
                    seen.add(text)
                    builder.add(text, f"section[{section_index}]/{label}/p[{i}]", content=False)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if any(n in archive.namelist() for n in ("word/footnotes.xml", "word/endnotes.xml")):
            builder.warn("notes_unread", "Ескертпе/соңғы сілтемелердің мазмұны жеке өңделмеді.")
    if document.element.xpath(".//w:ins | .//w:del"):
        builder.warn("tracked_changes", "Құжатта қабылданбаған өзгерістер бар. Соңғы бекітілген редакцияны тексеріңіз.")


def parse_pdf(content, builder):
    if not content.startswith(b"%PDF-"):
        raise RejectedFile("invalid_pdf")
    reader = PdfReader(io.BytesIO(content), strict=False)
    if reader.is_encrypted:
        raise RejectedFile("encrypted_pdf")
    for i, page in enumerate(reader.pages, 1):
        if i > 500 or builder.stopped:
            builder.warn("page_limit", "PDF-тің қалған беттері өңдеу шегіне байланысты оқылмады.")
            break
        try:
            resources = page.get("/Resources", {})
            resources = resources.get_object() if hasattr(resources, "get_object") else resources
            xobjects = resources.get("/XObject", {})
            xobjects = xobjects.get_object() if hasattr(xobjects, "get_object") else xobjects
            if any(obj.get_object().get("/Subtype") in ("/Image", "/Form") for obj in xobjects.values()):
                builder.warn("graphics_unread", f"PDF {i}-бетінде графикалық мазмұн бар; ол толық талданбады.")
            text = page.extract_text() or ""
            if not text.strip():
                builder.warn("page_without_text", f"PDF {i}-бетінен мәтін табылмады. OCR/мәтіндік балама қажет.")
            for j, paragraph in enumerate(text.splitlines(keepends=True), 1):
                builder.paragraph(paragraph, f"page[{i}]/line[{j}]", page=i)
        except Exception:
            builder.warn("page_parse_failed", f"PDF {i}-бетін оқу мүмкін болмады.")


def parse_xlsx(content, builder):
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in workbook:
            builder.headings.clear()
            builder.clauses.clear()
            # Read actual coordinates, not the untrusted cached <dimension> field.
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                xml = SafeXML.fromstring(archive.read(sheet._worksheet_path))
            coordinates = [
                coordinate_to_tuple(c.attrib["r"]) for c in xml.iter() if c.tag.endswith("}c") and "r" in c.attrib
            ]
            max_row = max((r for r, _ in coordinates), default=0)
            max_col = max((c for _, c in coordinates), default=0)
            if max_row > 20_000 or max_col > 256:
                builder.warn("sheet_limit", f"{sheet.title}: 20000 жол/256 баған шегінен тыс ұяшықтар оқылмады.")
            # Reset stale worksheet dimensions; explicit maxima still bound the work.
            sheet.reset_dimensions()
            column_context = {}
            for row in sheet.iter_rows(max_row=min(max_row, 20_000), max_col=min(max_col, 256)):
                if builder.stopped:
                    break
                row_context = []
                for cell in row:
                    if cell.value is None:
                        continue
                    formula = cell.data_type == "f"
                    count_before = len(builder.spans)
                    builder.add(
                        str(cell.value),
                        f"sheet[{sheet.title}]/{cell.coordinate}",
                        sheet=sheet.title,
                        cells=cell.coordinate,
                        content=not formula,
                    )
                    if len(builder.spans) > count_before:
                        span = builder.spans[-1]
                        span["context_span_ids"] = list(
                            dict.fromkeys(span["context_span_ids"] + column_context.get(cell.column, []) + row_context)
                        )
                        if not formula:
                            row_context.append(span["id"])
                            if cell.column not in column_context:
                                column_context[cell.column] = [span["id"]]
                    if formula:
                        builder.warn(
                            "formula_not_evaluated", f"{sheet.title}!{cell.coordinate}: формула орындалған жоқ."
                        )
    finally:
        workbook.close()


def parse_document(content: bytes, filename: str, kind: str, document_id: str, version: str, settings: Settings):
    builder = Builder(document_id, kind, settings)
    try:
        if kind in ("docx", "xlsx"):
            for flag in sorted(inspect_archive(content, kind, settings.max_archive_bytes)):
                builder.warn(flag, "Графикалық не сыртқы кірістірілген мазмұн оқылмады; мәтіндік баламасын тексеріңіз.")
        {"docx": parse_docx, "pdf": parse_pdf, "xlsx": parse_xlsx}[kind](content, builder)
    except RejectedFile:
        raise
    except Exception as exc:
        raise RejectedFile("parse_failed") from exc
    if not any(s["is_content"] for s in builder.spans):
        builder.warn("no_content", "Талдауға жарамды мәтін табылмады.")
    excerpt = "\n".join(s["raw_text"] for s in builder.spans[:15])
    lowered = excerpt.lower()
    types = [
        ("appendix", ("приложение", "қосымша")),
        ("order", ("приказ", "бұйрық")),
        ("job_description", ("должностная", "лауазымдық")),
        ("regulation", ("положение", "ереже")),
        ("structure", ("структура", "құрылым")),
    ]
    document_type = next((key for key, words in types if any(w in lowered for w in words)), "unknown")
    edition = re.search(r"(?:редакци[яи]|ред\.|редакция)\s*№?\s*\d+|\b\d{2}\.\d{2}\.\d{4}\b", excerpt, re.I)
    status = (
        "failed" if not any(s["is_content"] for s in builder.spans) else ("partial" if builder.warnings else "complete")
    )
    result = {
        "id": document_id,
        "version": version,
        "filename": filename,
        "title": next((s["raw_text"].strip()[:240] for s in builder.spans if s["raw_text"].strip()), filename),
        "document_type": document_type,
        "edition": edition.group(0) if edition else None,
        "source_sha256": hashlib.sha256(content).hexdigest(),
        "parse_status": status,
        "span_count": len(builder.spans),
        "warnings": builder.warnings,
        "spans": builder.spans,
    }
    return ParsedDocument.model_validate(result).model_dump()
